"""Build DANA4830_Final_Report.docx from analysis outputs."""
from __future__ import annotations

import json
import pathlib

from docx import Document
from docx.shared import Inches

ROOT = pathlib.Path(__file__).resolve().parent.parent
OUT = ROOT / "04_Analysis"
FIG = OUT / "figures"


def _add_md_paragraphs(doc: Document, path: pathlib.Path) -> None:
    if not path.exists():
        return
    for line in path.read_text(encoding="utf-8").splitlines():
        if line.startswith("# "):
            doc.add_heading(line[2:].strip(), level=1)
        elif line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=2)
        elif line.startswith("|") and "---" not in line:
            doc.add_paragraph(line)
        elif line.strip():
            doc.add_paragraph(line)


def main() -> None:
    report = json.loads((OUT / "report_data.json").read_text(encoding="utf-8"))
    doc = Document()
    doc.add_heading("DANA 4830 Final Exam Report", level=0)
    doc.add_paragraph("Student: Leo Sanchez")
    doc.add_paragraph(
        "This report summarizes Tasks 1 and 2. Numbers come from executed notebooks "
        "in 03_Codes/ and from 04_Analysis/report_data.json."
    )

    doc.add_heading("Task 1.1 — Comparison of approaches", level=1)
    t11 = OUT / "task1_1_comparison.md"
    if t11.exists():
        body = []
        for line in t11.read_text(encoding="utf-8").splitlines():
            if line.startswith("#"):
                continue
            body.append(line)
        doc.add_paragraph("\n".join(body).strip())
    else:
        doc.add_paragraph("See 04_Analysis/task1_1_comparison.md.")

    doc.add_heading("Task 1.2 — Pipeline diagrams", level=1)
    for name, caption in [
        ("diagram1_rosati_dge_pipeline.png", "Diagram 1: Rosati DGE pipeline (limma executed; RNA-seq tools shown as context)"),
        ("diagram2_midterm1_ml_pipeline.png", "Diagram 2: Midterm 1 / Syed ML pipeline (syed_pipeline)"),
        ("diagram_tmm_normalization.png", "TMM sketch (RNA-seq library-size scaling; not applied to GSE75214)"),
    ]:
        path = FIG / name
        if path.exists():
            doc.add_paragraph(caption)
            doc.add_picture(str(path), width=Inches(6.0))

    doc.add_heading("Task 2.1 — GSE75214 (IBD)", level=1)
    doc.add_paragraph(
        f"DEGs (FDR < 0.05, |log2FC| ≥ 1.0): {report['t21_deg']} "
        f"(up {report['t21_up']}, down {report['t21_down']}; "
        f"{report.get('t21_deg15', '—')} at |log2FC| ≥ 1.5). "
        f"Top genes: {', '.join(str(g[0]) for g in report['t21_top'][:5])}."
    )
    doc.add_paragraph(
        f"ML panel size: {report['t21_ml_n']} genes. "
        f"Midterm 1 Table-3 panel: 8 genes. "
        f"Overlap ML vs MT1: {report['t21_cmp'][1][1]}. "
        f"MT1 genes present on GSE75214: {report['t21_cmp'][-1][1]} of "
        f"{report['t21_cmp'][-1][2]}."
    )
    doc.add_paragraph(
        "limma is used because GSE75214 is microarray. SMOTE is applied on training "
        "data / inside CV folds. Survival analysis is not applicable (no follow-up). "
        "AUC on this cohort is optimistic if feature selection used the same samples. "
        "Zero overlap with the breast-cancer panel is limited transferability, not a "
        "failed Midterm 1."
    )

    doc.add_heading("Task 2.2 — Cervical merge", level=1)
    doc.add_paragraph(
        f"Merged dataset: {report['t22_samples']} samples, {report['t22_genes']} genes. "
        f"DEGs (FDR < 0.05, |log2FC| ≥ 1.5): {report['t22_deg']} "
        f"(up {report['t22_up']}, down {report['t22_down']}). "
        f"Wang hub genes recovered: {report['t22_wang_rec']}/10. "
        f"ML panel: {report['t22_ml_n']} genes; AURKA in panel: {report['t22_aurka']}."
    )
    doc.add_paragraph(
        "Study-wise centering is not ComBat. Wang intersected per-study DEGs; this "
        "exam merges samples first. 10/10 hubs is an external benchmark, not an exact replica."
    )

    doc.add_heading("Glossary", level=1)
    _add_md_paragraphs(doc, OUT / "glossary.md")

    doc.add_heading("Defense checklist", level=1)
    _add_md_paragraphs(doc, OUT / "DEFENSE_CHECKLIST.md")

    out_path = ROOT / "DANA4830_Final_Report.docx"
    doc.save(out_path)
    t11_docx = ROOT / "Task_1.1.docx"
    t11doc = Document()
    t11doc.add_heading("Task 1.1", level=0)
    t11doc.add_paragraph("Leo Sanchez — DANA 4830")
    if t11.exists():
        t11doc.add_paragraph(
            "\n".join(
                ln for ln in t11.read_text(encoding="utf-8").splitlines() if not ln.startswith("#")
            ).strip()
        )
    t11doc.save(t11_docx)
    print(f"Wrote {out_path}")
    print(f"Wrote {t11_docx}")


if __name__ == "__main__":
    main()
